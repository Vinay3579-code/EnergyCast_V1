from pathlib import Path
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parent.parent
RAW_DIR = PROJECT_ROOT / "data" / "raw" / "grid_india"


FILES_TO_CHECK = [
    RAW_DIR / "April 2022.xlsx",                 # 5-minute regime
    RAW_DIR / "April 2023.xlsx",                 # 10-second regime
    RAW_DIR / "January 2024- June 2025.xlsx",    # hourly regime
]


def inspect_workbook(path):
    print("\n")
    print("=" * 100)
    print(f"FILE: {path.name}")
    print(f"SIZE: {path.stat().st_size / (1024**2):.2f} MB")
    print("=" * 100)

    wb = load_workbook(
        path,
        read_only=True,
        data_only=True
    )

    print("\nSHEETS:")
    for sheet_name in wb.sheetnames:
        print(f" - {sheet_name}")

    for sheet_name in wb.sheetnames:

        ws = wb[sheet_name]

        print("\n" + "-" * 100)
        print(f"SHEET: {sheet_name}")
        print(f"Rows: {ws.max_row}")
        print(f"Columns: {ws.max_column}")
        print("-" * 100)

        print("\nFIRST 10 ROWS:")

        for row_number, row in enumerate(
            ws.iter_rows(
                min_row=1,
                max_row=min(10, ws.max_row),
                values_only=True
            ),
            start=1
        ):

            # Do not print hundreds of empty cells
            values = list(row[:30])

            print(
                f"ROW {row_number}:",
                values
            )

    wb.close()


for file in FILES_TO_CHECK:

    if not file.exists():
        print(f"\nERROR - FILE NOT FOUND: {file}")
        continue

    inspect_workbook(file)
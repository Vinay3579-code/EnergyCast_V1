from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


# ============================================================
# PATHS
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parent.parent

RAW_DIR = PROJECT_ROOT / "data" / "raw" / "grid_india"
OUTPUT_DIR = PROJECT_ROOT / "data" / "processed" / "grid_india"

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

MODEL_START = pd.Timestamp("2022-04-01 00:00:00")


# ============================================================
# HELPERS
# ============================================================

def get_sheet_names(path):
    wb = load_workbook(
        path,
        read_only=True,
        data_only=True
    )

    names = wb.sheetnames
    wb.close()

    return names


def infer_sampling_minutes(timestamp):
    timestamp = (
        pd.to_datetime(timestamp, errors="coerce")
        .dropna()
        .sort_values()
    )

    diff = (
        timestamp.diff()
        .dropna()
        .dt.total_seconds()
    )

    diff = diff[diff > 0]

    if len(diff) == 0:
        return np.nan

    return diff.median() / 60


# ============================================================
# OLD SCADA FORMAT
# ============================================================

def process_scada_file(path):

    # First two rows contain SCADA tag names
    # and human-readable signal names.
    header = pd.read_excel(
        path,
        sheet_name="Sheet1",
        header=None,
        nrows=2,
        engine="openpyxl"
    )

    readable_headers = [
        None if pd.isna(x) else str(x).strip()
        for x in header.iloc[1]
    ]

    def find_column(name):
        for i, value in enumerate(readable_headers):
            if value == name:
                return i

        return None


    demand_col = find_column("NLDC_DEMAND|P")
    wind_col = find_column("ALL_INDIA_WIND|P")

    # Prefer cleaned solar column
    solar_col = find_column("Solar")

    if solar_col is None:
        solar_col = find_column("ALL_IND_SOLAR|P")


    if demand_col is None:
        raise ValueError("NLDC_DEMAND|P not found")

    if wind_col is None:
        raise ValueError("ALL_INDIA_WIND|P not found")

    if solar_col is None:
        raise ValueError("Solar column not found")


    # Load only the columns we actually need.
    df = pd.read_excel(
        path,
        sheet_name="Sheet1",
        header=None,
        skiprows=2,
        usecols=[
            0,
            wind_col,
            solar_col,
            demand_col
        ],
        engine="openpyxl"
    )


    df.columns = [
        "timestamp",
        "wind_mw",
        "solar_mw",
        "demand_mw"
    ]


    # --------------------------------------------------------
    # TYPES
    # --------------------------------------------------------

    df["timestamp"] = pd.to_datetime(
        df["timestamp"],
        errors="coerce"
    )

    df = df.dropna(
        subset=["timestamp"]
    )


    for col in [
        "demand_mw",
        "wind_mw",
        "solar_mw"
    ]:

        df[col] = pd.to_numeric(
            df[col],
            errors="coerce"
        )


    sampling_minutes = infer_sampling_minutes(
        df["timestamp"]
    )


    print(
        f"    Detected sampling: "
        f"{sampling_minutes:.4f} minutes"
    )


    # --------------------------------------------------------
    # HIGH-FREQUENCY -> HOURLY
    # --------------------------------------------------------

    df = (
        df
        .set_index("timestamp")
        .sort_index()
    )


    values = (
        df[
            [
                "demand_mw",
                "wind_mw",
                "solar_mw"
            ]
        ]
        .resample("1h")
        .mean()
    )


    counts = (
        df[
            [
                "demand_mw",
                "wind_mw",
                "solar_mw"
            ]
        ]
        .resample("1h")
        .count()
    )


    values["demand_samples"] = counts["demand_mw"]
    values["wind_samples"] = counts["wind_mw"]
    values["solar_samples"] = counts["solar_mw"]


    values = values.reset_index()


    values["source_file"] = path.name

    values["source_format"] = "SCADA"

    values["raw_sampling_minutes"] = sampling_minutes


    return values


# ============================================================
# NEW HOURLY FORMAT
# ============================================================

def process_hourly_report(path):

    df = pd.read_excel(
        path,
        sheet_name="Report",
        engine="openpyxl"
    )


    required_columns = [
        "Timestamp",
        "Demand (MW)",
        "Wind (MW)",
        "Solar (MW)"
    ]


    missing = [
        col
        for col in required_columns
        if col not in df.columns
    ]


    if missing:
        raise ValueError(
            f"Missing columns: {missing}"
        )


    df = df[
        required_columns
    ].copy()


    df = df.rename(
        columns={
            "Timestamp": "timestamp",
            "Demand (MW)": "demand_mw",
            "Wind (MW)": "wind_mw",
            "Solar (MW)": "solar_mw"
        }
    )


    df["timestamp"] = pd.to_datetime(
        df["timestamp"],
        errors="coerce",
        dayfirst=True
    )


    df = df.dropna(
        subset=["timestamp"]
    )


    for col in [
        "demand_mw",
        "wind_mw",
        "solar_mw"
    ]:

        df[col] = pd.to_numeric(
            df[col],
            errors="coerce"
        )


    df["demand_samples"] = 1
    df["wind_samples"] = 1
    df["solar_samples"] = 1

    df["source_file"] = path.name

    df["source_format"] = "HOURLY"

    df["raw_sampling_minutes"] = 60


    return df


# ============================================================
# MAIN
# ============================================================

files = sorted(
    RAW_DIR.glob("*.xlsx")
)


print("=" * 75)
print("ENERGYCAST V1 - GRID INDIA PREPROCESSOR")
print("=" * 75)

print(f"Raw folder: {RAW_DIR}")
print(f"Files found: {len(files)}")


if len(files) == 0:
    raise FileNotFoundError(
        "No Excel files found."
    )


frames = []

validation = []

failed = []


# ============================================================
# PROCESS FILES
# ============================================================

for i, file in enumerate(files, start=1):

    print()
    print(
        f"[{i:02d}/{len(files)}] "
        f"{file.name}"
    )


    try:

        sheets = get_sheet_names(file)


        if "Report" in sheets:

            print("    Format: Hourly Report")

            df = process_hourly_report(
                file
            )


        elif "Sheet1" in sheets:

            print("    Format: SCADA")

            df = process_scada_file(
                file
            )


        else:

            raise ValueError(
                f"Unknown workbook structure: {sheets}"
            )


        start = df["timestamp"].min()

        end = df["timestamp"].max()


        print(
            f"    Range: {start} -> {end}"
        )

        print(
            f"    Hourly rows: {len(df):,}"
        )


        validation.append({
            "file": file.name,
            "start": start,
            "end": end,
            "hourly_rows": len(df),
            "missing_demand":
                df["demand_mw"].isna().sum(),
            "missing_wind":
                df["wind_mw"].isna().sum(),
            "missing_solar":
                df["solar_mw"].isna().sum()
        })


        frames.append(df)


    except Exception as e:

        print(
            f"    ERROR: {e}"
        )

        failed.append({
            "file": file.name,
            "error": str(e)
        })


# ============================================================
# VALIDATION OUTPUT
# ============================================================

validation_df = pd.DataFrame(
    validation
)


validation_df.to_csv(
    OUTPUT_DIR /
    "grid_india_file_validation.csv",
    index=False
)


if failed:

    pd.DataFrame(
        failed
    ).to_csv(
        OUTPUT_DIR /
        "grid_india_failed_files.csv",
        index=False
    )

    raise RuntimeError(
        f"{len(failed)} files failed. "
        "Check grid_india_failed_files.csv"
    )


# ============================================================
# MERGE
# ============================================================

combined = pd.concat(
    frames,
    ignore_index=True
)


combined = combined[
    combined["timestamp"] >= MODEL_START
].copy()


combined = combined.sort_values(
    "timestamp"
).reset_index(drop=True)


# ============================================================
# DUPLICATES
# ============================================================

duplicate_mask = (
    combined["timestamp"]
    .duplicated(keep=False)
)


duplicates = combined[
    duplicate_mask
]


print()
print(
    "Duplicate timestamp rows:",
    len(duplicates)
)


if len(duplicates) > 0:

    duplicates.to_csv(
        OUTPUT_DIR /
        "grid_india_duplicate_timestamps.csv",
        index=False
    )


# ============================================================
# MISSING TIMESTAMPS
# ============================================================

start = combined["timestamp"].min()

end = combined["timestamp"].max()


expected = pd.date_range(
    start=start,
    end=end,
    freq="1h"
)


actual = pd.DatetimeIndex(
    combined["timestamp"].unique()
)


missing = expected.difference(
    actual
)


print(
    "Missing hourly timestamps:",
    len(missing)
)


if len(missing) > 0:

    pd.DataFrame({
        "missing_timestamp": missing
    }).to_csv(
        OUTPUT_DIR /
        "grid_india_missing_timestamps.csv",
        index=False
    )


# ============================================================
# SAVE FINAL DATA
# ============================================================

FINAL_COLUMNS = [
    "timestamp",
    "demand_mw",
    "wind_mw",
    "solar_mw",
    "demand_samples",
    "wind_samples",
    "solar_samples",
    "source_file",
    "source_format",
    "raw_sampling_minutes"
]


combined = combined[
    FINAL_COLUMNS
]


output_file = (
    OUTPUT_DIR /
    "grid_india_hourly_2022_2025.csv"
)


combined.to_csv(
    output_file,
    index=False
)


# ============================================================
# FINAL SUMMARY
# ============================================================

print()
print("=" * 75)
print("FINAL QUALITY CHECK")
print("=" * 75)

print(
    "Rows:",
    len(combined)
)

print(
    "Start:",
    combined["timestamp"].min()
)

print(
    "End:",
    combined["timestamp"].max()
)

print(
    "Missing timestamps:",
    len(missing)
)

print(
    "Missing demand:",
    combined["demand_mw"].isna().sum()
)

print(
    "Missing wind:",
    combined["wind_mw"].isna().sum()
)

print(
    "Missing solar:",
    combined["solar_mw"].isna().sum()
)

print()
print(
    "Saved:",
    output_file
)